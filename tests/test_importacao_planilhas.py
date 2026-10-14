import io

import openpyxl

from app.models.cliente import Cliente
from app.utils import planilhas
from tests.conftest import dados_cliente

CAMPOS_LINHA_PADRAO = {
    "tipo_pessoa": "F",
    "nome": "Joana Prado",
    "nome_fantasia": "",
    "documento": "529.982.247-25",
    "email": "joana@exemplo.com",
    "telefone": "(11) 91234-5678",
    "cep": "",
    "logradouro": "",
    "numero": "",
    "complemento": "",
    "bairro": "",
    "cidade": "",
    "uf": "",
    "observacoes": "",
    "contato_nome": "Carlos",
    "contato_cargo": "",
    "contato_email": "",
    "contato_telefone": "",
}


def _linha(**alteracoes):
    dados = dict(CAMPOS_LINHA_PADRAO, **alteracoes)
    return [dados[coluna] for coluna in planilhas.COLUNAS]


def _construir_csv(linhas):
    buffer = io.StringIO()
    buffer.write(",".join(planilhas.COLUNAS) + "\n")
    for linha in linhas:
        buffer.write(",".join(linha) + "\n")
    return io.BytesIO(buffer.getvalue().encode("utf-8"))


def _token_da_url(resposta):
    # A pre-visualizacao fica em /clientes/importar/<token>.
    return resposta.request.path.rsplit("/", 1)[-1]


def test_importa_planilha_csv_valida_com_contato(client_logado):
    arquivo = _construir_csv([_linha()])

    resposta = client_logado.post(
        "/clientes/importar",
        data={"arquivo": (arquivo, "clientes.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert resposta.status_code == 200
    assert "Pronto" in resposta.get_data(as_text=True)

    token = _token_da_url(resposta)
    resposta = client_logado.post(
        f"/clientes/importar/{token}/confirmar", follow_redirects=True
    )

    assert resposta.status_code == 200
    cliente = Cliente.query.filter_by(documento="52998224725").first()
    assert cliente is not None
    assert cliente.contatos[0].nome == "Carlos"
    assert cliente.contatos[0].principal is True


def test_linha_com_cpf_invalido_nao_e_importada(client_logado):
    arquivo = _construir_csv([_linha(documento="111.111.111-11")])

    resposta = client_logado.post(
        "/clientes/importar",
        data={"arquivo": (arquivo, "clientes.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "CPF invalido" in resposta.get_data(as_text=True)

    token = _token_da_url(resposta)
    client_logado.post(f"/clientes/importar/{token}/confirmar", follow_redirects=True)

    assert Cliente.query.count() == 0


def test_linha_com_documento_ja_cadastrado_nao_e_importada(client_logado):
    client_logado.post("/clientes/novo", data=dados_cliente(), follow_redirects=True)

    arquivo = _construir_csv(
        [_linha(documento="529.982.247-25", nome="Outro Nome", contato_nome="")]
    )
    resposta = client_logado.post(
        "/clientes/importar",
        data={"arquivo": (arquivo, "clientes.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "Ja cadastrado" in resposta.get_data(as_text=True)

    token = _token_da_url(resposta)
    client_logado.post(f"/clientes/importar/{token}/confirmar", follow_redirects=True)

    assert Cliente.query.count() == 1
    assert Cliente.query.first().nome == "Maria Silva"


def test_documento_repetido_na_propria_planilha_e_marcado_como_erro(client_logado):
    arquivo = _construir_csv(
        [
            _linha(documento="529.982.247-25", nome="Primeira"),
            _linha(documento="529.982.247-25", nome="Segunda"),
        ]
    )
    resposta = client_logado.post(
        "/clientes/importar",
        data={"arquivo": (arquivo, "clientes.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    corpo = resposta.get_data(as_text=True)
    assert "repetido" in corpo

    token = _token_da_url(resposta)
    client_logado.post(f"/clientes/importar/{token}/confirmar", follow_redirects=True)

    # So a primeira ocorrencia fica marcada como "ok"; nenhuma das duas com o
    # mesmo documento pode virar dois clientes.
    assert Cliente.query.filter_by(documento="52998224725").count() == 1


def test_exportar_clientes_csv_contem_o_cliente_cadastrado(client_logado):
    client_logado.post("/clientes/novo", data=dados_cliente(), follow_redirects=True)

    resposta = client_logado.get("/clientes/exportar?formato=csv")

    assert resposta.status_code == 200
    assert resposta.mimetype == "text/csv"
    assert "Maria Silva" in resposta.get_data(as_text=True)


def test_exportar_clientes_xlsx_contem_o_cliente_cadastrado(client_logado):
    client_logado.post("/clientes/novo", data=dados_cliente(), follow_redirects=True)

    resposta = client_logado.get("/clientes/exportar?formato=xlsx")
    assert resposta.status_code == 200

    livro = openpyxl.load_workbook(io.BytesIO(resposta.data))
    valores = list(livro.active.iter_rows(values_only=True))

    assert any(linha[0] == "Maria Silva" for linha in valores[1:])


def test_exportar_respeita_filtro_de_situacao(client_logado):
    client_logado.post("/clientes/novo", data=dados_cliente(), follow_redirects=True)
    cliente = Cliente.query.first()
    client_logado.post(f"/clientes/{cliente.id}/situacao", follow_redirects=True)

    resposta = client_logado.get("/clientes/exportar?formato=csv&situacao=ativos")

    assert "Maria Silva" not in resposta.get_data(as_text=True)


def _enviar(client, conteudo, nome):
    return client.post(
        "/clientes/importar",
        data={"arquivo": (io.BytesIO(conteudo), nome)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )


def test_baixa_planilha_modelo(client_logado):
    resposta = client_logado.get("/clientes/importar/modelo")

    planilha = openpyxl.load_workbook(io.BytesIO(resposta.data)).active
    assert [c.value for c in planilha[1]] == planilhas.COLUNAS


def test_xlsx_corrompido_mostra_aviso_sem_erro_500(client_logado):
    resposta = _enviar(client_logado, b"nao e um zip", "ruim.xlsx")

    assert resposta.status_code == 200
    assert "corrompido" in resposta.get_data(as_text=True)


def test_csv_com_codificacao_invalida_mostra_aviso(client_logado):
    resposta = _enviar(client_logado, b"\xff\xfe\xfa", "ruim.csv")

    assert resposta.status_code == 200
    assert Cliente.query.count() == 0


def test_planilha_so_com_cabecalho_avisa_que_esta_vazia(client_logado):
    resposta = _enviar(client_logado, (",".join(planilhas.COLUNAS) + "\n").encode(), "v.csv")

    assert "A planilha nao tem nenhuma linha para importar." in resposta.get_data(as_text=True)


def test_token_inexistente_volta_para_o_envio(client_logado):
    token = "0" * 32
    for metodo, url in [
        ("get", f"/clientes/importar/{token}"),
        ("get", f"/clientes/importar/{token}/erros"),
        ("post", f"/clientes/importar/{token}/confirmar"),
    ]:
        resposta = getattr(client_logado, metodo)(url)
        assert resposta.status_code == 302, url
        assert resposta.headers["Location"] == "/clientes/importar", url


def test_relatorio_de_erros_traz_a_linha_e_o_motivo(client_logado):
    previa = client_logado.post(
        "/clientes/importar",
        data={"arquivo": (_construir_csv([_linha(documento="111.111.111-11")]), "c.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    resposta = client_logado.get(f"/clientes/importar/{_token_da_url(previa)}/erros")

    texto = resposta.data.decode("utf-8-sig")
    assert resposta.mimetype == "text/csv"
    assert "CPF invalido." in texto


def test_documento_cadastrado_entre_previa_e_confirmacao_nao_duplica(client_logado):
    previa = client_logado.post(
        "/clientes/importar",
        data={"arquivo": (_construir_csv([_linha()]), "c.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    client_logado.post("/clientes/novo", data=dados_cliente(documento="529.982.247-25"))

    client_logado.post(
        f"/clientes/importar/{_token_da_url(previa)}/confirmar", follow_redirects=True
    )

    assert Cliente.query.count() == 1


def test_confirmacao_so_vale_uma_vez(client_logado):
    previa = client_logado.post(
        "/clientes/importar",
        data={"arquivo": (_construir_csv([_linha()]), "c.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    url = f"/clientes/importar/{_token_da_url(previa)}/confirmar"

    client_logado.post(url)
    segunda = client_logado.post(url)

    assert segunda.headers["Location"] == "/clientes/importar"
    assert Cliente.query.count() == 1
