"""Testes unitarios das funcoes de app/utils/planilhas.py."""

import csv
import io

import openpyxl
import pytest

from app.utils import planilhas


def _xlsx(*linhas):
    livro = openpyxl.Workbook()
    for linha in linhas:
        livro.active.append(linha)
    buffer = io.BytesIO()
    livro.save(buffer)
    buffer.seek(0)
    return buffer


def _linha(**campos):
    return {coluna: "" for coluna in planilhas.COLUNAS} | campos


def test_le_xlsx_normaliza_cabecalho_e_ignora_extras_e_vazias():
    arquivo = _xlsx(
        ["NOME", "Documento", "coluna_desconhecida"],
        ["  Ana  ", 52998224725, "x"],
        [None, None, None],
    )

    linhas = planilhas.ler_planilha(arquivo, "clientes.XLSX")

    assert len(linhas) == 1
    assert linhas[0]["nome"] == "Ana"
    assert linhas[0]["documento"] == "52998224725"
    assert linhas[0]["email"] == ""
    assert "coluna_desconhecida" not in linhas[0]


def test_xlsx_corrompido_vira_erro_de_leitura():
    with pytest.raises(ValueError, match="corrompido"):
        planilhas.ler_planilha(io.BytesIO(b"nao e um zip"), "ruim.xlsx")


def test_recusa_extensao_nao_suportada():
    with pytest.raises(ValueError, match="nao suportado"):
        planilhas.ler_planilha(io.BytesIO(b""), "clientes.txt")


def test_csv_vazio_devolve_lista_vazia():
    assert planilhas.ler_planilha(io.BytesIO(b""), "vazio.csv") == []


def test_planilha_modelo_tem_cabecalhos_e_exemplo_valido(app):
    planilha = openpyxl.load_workbook(planilhas.gerar_modelo_importacao()).active
    linhas = [list(linha) for linha in planilha.iter_rows(values_only=True)]

    assert linhas[0] == planilhas.COLUNAS
    exemplo = dict(zip(planilhas.COLUNAS, linhas[1]))
    assert planilhas.validar_linhas([exemplo])[0]["status"] == "ok"


@pytest.mark.parametrize(
    "campos, erro",
    [
        ({"tipo_pessoa": "X", "nome": "A", "documento": "1"}, "tipo_pessoa deve ser F ou J."),
        ({"tipo_pessoa": "F", "documento": "529.982.247-25"}, "nome e obrigatorio."),
        ({"tipo_pessoa": "F", "nome": "A"}, "documento e obrigatorio."),
        ({"tipo_pessoa": "J", "nome": "A", "documento": "11.111.111/0001-11"}, "CNPJ invalido."),
        (
            {"tipo_pessoa": "F", "nome": "A", "documento": "529.982.247-25", "email": "sem-arroba"},
            "e-mail invalido.",
        ),
        (
            {"tipo_pessoa": "F", "nome": "A", "documento": "529.982.247-25", "contato_email": "x@"},
            "e-mail do contato invalido.",
        ),
    ],
)
def test_validacao_aponta_cada_erro(app, campos, erro):
    resultado = planilhas.validar_linhas([_linha(**campos)])[0]

    assert resultado["status"] == "erro"
    assert erro in resultado["erros"]


def test_validacao_normaliza_campos(app):
    linha = _linha(
        tipo_pessoa="f", nome=" Ana ", documento="529.982.247-25", email="ANA@X.COM", uf="sp"
    )
    dados = planilhas.validar_linhas([linha])[0]["dados"]

    assert (dados["tipo_pessoa"], dados["nome"], dados["documento"]) == ("F", "Ana", "52998224725")
    assert (dados["email"], dados["uf"]) == ("ana@x.com", "SP")


def test_relatorio_lista_so_linhas_rejeitadas_e_neutraliza_formulas(app):
    pendente = {
        "linhas": planilhas.validar_linhas(
            [
                _linha(tipo_pessoa="F", nome="Valida", documento="529.982.247-25"),
                _linha(tipo_pessoa="X", nome="=HYPERLINK(\"x\")", documento="1"),
            ]
        )
    }

    texto = planilhas.gerar_relatorio_erros(pendente).getvalue().decode("utf-8-sig")
    registros = list(csv.reader(io.StringIO(texto)))

    assert len(registros) == 2  # cabecalho + a linha com erro
    assert registros[1][1] == "erro"
    assert "tipo_pessoa deve ser F ou J." in registros[1][2]
    assert "'=HYPERLINK" in texto


@pytest.mark.parametrize("token", ["", "../../etc/passwd", "ABC", "0" * 31])
def test_token_com_formato_invalido_e_recusado(token):
    assert planilhas.carregar_pendente(token) is None
    planilhas.descartar_pendente(token)  # nao deve levantar erro


def test_pendente_salvo_carrega_e_descarta(app):
    token = planilhas.salvar_pendente([], "a.csv")

    assert planilhas.carregar_pendente(token)["arquivo"] == "a.csv"
    planilhas.descartar_pendente(token)
    assert planilhas.carregar_pendente(token) is None
