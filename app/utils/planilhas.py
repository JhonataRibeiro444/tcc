"""Importacao e exportacao de clientes e contatos via planilha (.xlsx/.csv).

Uma linha da planilha corresponde a um Cliente (PF ou PJ) e, opcionalmente,
ao contato principal dele. A importacao acontece em duas etapas: a planilha
enviada e lida e validada, o resultado fica guardado temporariamente sob um
token, e so e gravado no banco quando o usuario confirma a pre-visualizacao.

Duplicidade de CPF/CNPJ nunca sobrescreve um cliente existente: a linha e
apenas marcada como "duplicado" e nao entra na importacao, seguindo a mesma
regra usada no cadastro manual (ClienteForm.validate_documento).
"""

import csv
import io
import json
import os
import re
import tempfile
import uuid
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook

from app.models.cliente import Cliente
from app.utils.validators import apenas_digitos, formatar_documento, validar_documento

COLUNAS = [
    "tipo_pessoa",
    "nome",
    "nome_fantasia",
    "documento",
    "email",
    "telefone",
    "cep",
    "logradouro",
    "numero",
    "complemento",
    "bairro",
    "cidade",
    "uf",
    "observacoes",
    "contato_nome",
    "contato_cargo",
    "contato_email",
    "contato_telefone",
]

LINHA_EXEMPLO = [
    "F", "Maria Silva", "", "529.982.247-25", "maria@exemplo.com",
    "(11) 91234-5678", "01001-000", "Praca da Se", "100", "", "Se",
    "Sao Paulo", "SP", "",
    "Joao Souza", "Comprador", "joao@exemplo.com", "(11) 98888-7777",
]

COLUNAS_EXPORTACAO = [
    ("Nome / Razao social", lambda c: c.nome),
    ("Tipo", lambda c: c.tipo_pessoa_label),
    ("CPF / CNPJ", lambda c: c.documento_formatado),
    ("E-mail", lambda c: c.email or ""),
    ("Telefone", lambda c: c.telefone_formatado),
    ("Cidade", lambda c: c.cidade or ""),
    ("UF", lambda c: c.uf or ""),
    ("Situacao", lambda c: "Ativo" if c.ativo else "Inativo"),
    ("Contato principal", lambda c: c.contato_principal.nome if c.contato_principal else ""),
    (
        "E-mail do contato",
        lambda c: (c.contato_principal.email or "") if c.contato_principal else "",
    ),
    (
        "Telefone do contato",
        lambda c: c.contato_principal.telefone_formatado if c.contato_principal else "",
    ),
]

_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_TOKEN_RE = re.compile(r"^[0-9a-f]{32}$")
_DIR_TEMP = os.path.join(tempfile.gettempdir(), "crm_importacoes")


def _linha_vazia(valores):
    return all((valor is None or str(valor).strip() == "") for valor in valores)


def _sanitizar_celula_csv(valor):
    """Evita "CSV injection": prefixa com apostrofo formulas que o Excel
    poderia executar ao abrir o arquivo (celula comecando com =, +, - ou @)."""
    texto = "" if valor is None else str(valor)
    if texto[:1] in ("=", "+", "-", "@"):
        return "'" + texto
    return texto


# --------------------------------------------------------------------------
# Planilha modelo
# --------------------------------------------------------------------------


def gerar_modelo_importacao():
    """Devolve os bytes de uma planilha .xlsx modelo para a importacao."""
    livro = Workbook()
    planilha = livro.active
    planilha.title = "clientes"
    planilha.append(COLUNAS)
    planilha.append(LINHA_EXEMPLO)

    for coluna in planilha.columns:
        letra = coluna[0].column_letter
        maior = max(len(str(celula.value or "")) for celula in coluna)
        planilha.column_dimensions[letra].width = min(max(maior + 2, 10), 40)

    buffer = io.BytesIO()
    livro.save(buffer)
    buffer.seek(0)
    return buffer


# --------------------------------------------------------------------------
# Leitura do arquivo enviado
# --------------------------------------------------------------------------


def ler_planilha(stream, nome_arquivo):
    """Le um arquivo .xlsx ou .csv e devolve uma lista de dicionarios brutos.

    As chaves seguem `COLUNAS`. Uma coluna que nao aparecer no arquivo fica
    com string vazia; uma coluna desconhecida no arquivo e ignorada. Linhas
    totalmente vazias sao descartadas.
    """
    extensao = (nome_arquivo.rsplit(".", 1)[-1] if "." in nome_arquivo else "").lower()

    if extensao == "csv":
        texto = stream.read().decode("utf-8-sig")
        linhas_brutas = list(csv.reader(io.StringIO(texto)))
    elif extensao == "xlsx":
        try:
            livro = load_workbook(stream, data_only=True, read_only=True)
            planilha = livro.active
            linhas_brutas = [list(linha) for linha in planilha.iter_rows(values_only=True)]
        except Exception as erro:  # openpyxl falha com varios tipos (zip, xml, chave)
            raise ValueError("Nao foi possivel ler o arquivo .xlsx. Ele esta corrompido?") from erro
    else:
        raise ValueError("Formato de arquivo nao suportado. Envie um .xlsx ou .csv.")

    if not linhas_brutas:
        return []

    cabecalho = [str(celula or "").strip().lower() for celula in linhas_brutas[0]]

    linhas = []
    for bruta in linhas_brutas[1:]:
        if _linha_vazia(bruta):
            continue

        linha = {}
        for indice, chave in enumerate(cabecalho):
            if chave in COLUNAS:
                valor = bruta[indice] if indice < len(bruta) else ""
                linha[chave] = "" if valor is None else str(valor).strip()
        for coluna in COLUNAS:
            linha.setdefault(coluna, "")

        linhas.append(linha)

    return linhas


# --------------------------------------------------------------------------
# Validacao
# --------------------------------------------------------------------------


def validar_linhas(linhas):
    """Valida cada linha lida da planilha e classifica o resultado como
    "ok", "duplicado" (documento ja cadastrado) ou "erro".

    Reaproveita as mesmas regras do cadastro manual: CPF/CNPJ pelo modulo 11
    (app.utils.validators) e duplicidade por documento ja existente no banco.
    """
    processadas = []
    documentos_no_arquivo = {}

    for numero, linha in enumerate(linhas, start=2):  # a linha 1 e o cabecalho
        erros = []

        tipo_pessoa = (linha["tipo_pessoa"] or "").strip().upper()
        if tipo_pessoa not in ("F", "J"):
            erros.append("tipo_pessoa deve ser F ou J.")

        nome = (linha["nome"] or "").strip()
        if not nome:
            erros.append("nome e obrigatorio.")

        documento = apenas_digitos(linha["documento"])
        if not documento:
            erros.append("documento e obrigatorio.")
        elif tipo_pessoa in ("F", "J") and not validar_documento(documento, tipo_pessoa):
            rotulo = "CPF" if tipo_pessoa == "F" else "CNPJ"
            erros.append(f"{rotulo} invalido.")

        email = (linha["email"] or "").strip().lower()
        if email and not _EMAIL_RE.match(email):
            erros.append("e-mail invalido.")

        contato_nome = (linha["contato_nome"] or "").strip()
        contato_email = (linha["contato_email"] or "").strip().lower()
        if contato_email and not _EMAIL_RE.match(contato_email):
            erros.append("e-mail do contato invalido.")

        if not erros and documento:
            if documento in documentos_no_arquivo:
                erros.append(
                    "documento repetido na propria planilha "
                    f"(linha {documentos_no_arquivo[documento]})."
                )
            else:
                documentos_no_arquivo[documento] = numero

        status = "erro" if erros else "ok"
        if status == "ok" and Cliente.query.filter_by(documento=documento).first():
            status = "duplicado"

        processadas.append(
            {
                "linha": numero,
                "status": status,
                "erros": erros,
                "dados": {
                    "tipo_pessoa": tipo_pessoa,
                    "nome": nome,
                    "nome_fantasia": (linha["nome_fantasia"] or "").strip(),
                    "documento": documento,
                    "documento_formatado": formatar_documento(documento) if documento else "",
                    "email": email,
                    "telefone": apenas_digitos(linha["telefone"]),
                    "cep": apenas_digitos(linha["cep"]),
                    "logradouro": (linha["logradouro"] or "").strip(),
                    "numero": (linha["numero"] or "").strip(),
                    "complemento": (linha["complemento"] or "").strip(),
                    "bairro": (linha["bairro"] or "").strip(),
                    "cidade": (linha["cidade"] or "").strip(),
                    "uf": (linha["uf"] or "").strip().upper(),
                    "observacoes": (linha["observacoes"] or "").strip(),
                    "contato_nome": contato_nome,
                    "contato_cargo": (linha["contato_cargo"] or "").strip(),
                    "contato_email": contato_email,
                    "contato_telefone": apenas_digitos(linha["contato_telefone"]),
                },
            }
        )

    return processadas


def linhas_por_status(pendente):
    """Separa as linhas de uma importacao pendente pelos tres status possiveis."""
    ok, duplicados, erros = [], [], []
    for linha in pendente["linhas"]:
        if linha["status"] == "ok":
            ok.append(linha)
        elif linha["status"] == "duplicado":
            duplicados.append(linha)
        else:
            erros.append(linha)
    return ok, duplicados, erros


# --------------------------------------------------------------------------
# Importacao pendente (entre o upload e a confirmacao)
# --------------------------------------------------------------------------


def _caminho_pendente(token):
    # O token e sempre gerado por nos (uuid4 hex); validar aqui evita que um
    # valor arbitrario vindo da URL vire um caminho de arquivo.
    if not token or not _TOKEN_RE.match(token):
        return None
    return os.path.join(_DIR_TEMP, f"{token}.json")


def salvar_pendente(linhas_processadas, nome_arquivo_original):
    """Guarda o resultado da validacao em disco e devolve um token de acesso."""
    os.makedirs(_DIR_TEMP, exist_ok=True)
    token = uuid.uuid4().hex
    conteudo = {
        "arquivo": nome_arquivo_original,
        "criado_em": datetime.now(timezone.utc).isoformat(),
        "linhas": linhas_processadas,
    }
    with open(_caminho_pendente(token), "w", encoding="utf-8") as arquivo:
        json.dump(conteudo, arquivo)
    return token


def carregar_pendente(token):
    caminho = _caminho_pendente(token)
    if not caminho or not os.path.exists(caminho):
        return None
    with open(caminho, encoding="utf-8") as arquivo:
        return json.load(arquivo)


def descartar_pendente(token):
    caminho = _caminho_pendente(token)
    if caminho and os.path.exists(caminho):
        os.remove(caminho)


def gerar_relatorio_erros(pendente):
    """CSV com as linhas nao importadas (erro ou duplicado) e o motivo."""
    buffer = io.StringIO()
    escritor = csv.writer(buffer)
    escritor.writerow(["linha", "status", "motivo", *COLUNAS])

    for linha in pendente["linhas"]:
        if linha["status"] == "ok":
            continue
        motivo = (
            "Documento ja cadastrado."
            if linha["status"] == "duplicado"
            else "; ".join(linha["erros"])
        )
        escritor.writerow(
            [
                linha["linha"],
                linha["status"],
                _sanitizar_celula_csv(motivo),
                *[_sanitizar_celula_csv(linha["dados"].get(coluna, "")) for coluna in COLUNAS],
            ]
        )

    return io.BytesIO(buffer.getvalue().encode("utf-8-sig"))


# --------------------------------------------------------------------------
# Exportacao
# --------------------------------------------------------------------------


def exportar_clientes_xlsx(clientes):
    livro = Workbook()
    planilha = livro.active
    planilha.title = "clientes"
    planilha.append([titulo for titulo, _ in COLUNAS_EXPORTACAO])
    for cliente in clientes:
        planilha.append([extrator(cliente) for _, extrator in COLUNAS_EXPORTACAO])

    buffer = io.BytesIO()
    livro.save(buffer)
    buffer.seek(0)
    return buffer


def exportar_clientes_csv(clientes):
    buffer = io.StringIO()
    escritor = csv.writer(buffer)
    escritor.writerow([titulo for titulo, _ in COLUNAS_EXPORTACAO])
    for cliente in clientes:
        escritor.writerow(
            [_sanitizar_celula_csv(extrator(cliente)) for _, extrator in COLUNAS_EXPORTACAO]
        )
    return io.BytesIO(buffer.getvalue().encode("utf-8-sig"))
